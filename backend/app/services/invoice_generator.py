from datetime import datetime
from pathlib import Path
from typing import Dict, Any
from docx import Document as DocxDocument
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from ..config import settings


class InvoiceGenerator:
    """
    Generates a B2B Invoice Word document from an Income record.
    Matches the LSI invoice template.
    """

    def __init__(self):
        self.output_dir = Path(settings.OUTPUT_DIR)
        self.output_dir.mkdir(exist_ok=True)

    def generate_invoice_docx(self, data: Dict[str, Any]) -> str:
        doc = DocxDocument()

        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(2)
            section.right_margin = Cm(2)

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.space_after = Pt(0)

        # Seller header
        for text, bold, size in [
            ("Leapstack International Limited", True, 11),
            ("Unit 105A, 1/F., Building 5W, Hong Kong Science Park, Shatin, N.T. H.K.", False, 9),
            ("Tel: +852 3628 5039   Fax: +852 3105 5299", False, 9),
        ]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)

        doc.add_paragraph()

        # Title
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = title.add_run("I N V O I C E")
        r.bold = True
        r.font.size = Pt(16)
        title.paragraph_format.space_after = Pt(12)

        # Invoice header table: Invoice No | Date | Billed To
        header = doc.add_table(rows=1, cols=2)
        header.autofit = False

        left = header.rows[0].cells[0]
        left.text = ""
        p = left.paragraphs[0]
        p.add_run("INVOICE NUMBER\n").bold = True
        p.add_run(f"{data['invoice_number']}\n\n")
        p.add_run("DATE OF ISSUE\n").bold = True
        p.add_run(f"{data['date_issued']}\n\n")
        p.add_run("BILLED TO:\n").bold = True
        p.add_run(f"{data['customer_name']}\n")
        if data.get("customer_address"):
            p.add_run(f"{data['customer_address']}\n")

        right = header.rows[0].cells[1]
        right.text = ""

        doc.add_paragraph()

        # Line items table
        headers = ["No.", "Description", "Invoice No", "Machine Runs", "Rate (RM)", "Amount (RM)"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        hdr = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for idx, li in enumerate(data.get("line_items", []), start=1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = li.get("description", "")
            row[2].text = li.get("invoice_no") or ""
            row[3].text = str(li.get("machine_runs") or "")
            row[4].text = f"{li.get('rate', 0):.2f}"
            row[5].text = f"{li.get('amount', 0):,.2f}"

            for c in (0, 2, 3, 4, 5):
                row[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Total row
        total_row = table.add_row().cells
        total_row[0].merge(total_row[4])
        p = total_row[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("TOTAL (RM)")
        r.bold = True

        p = total_row[5].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"{data['total_rm']:,.2f}")
        r.bold = True

        # Optional HKD/USD lines
        for label, amount in [
            ("TOTAL (HKD)", data.get("total_hkd")),
            ("TOTAL (USD)", data.get("total_usd")),
        ]:
            if amount:
                r_row = table.add_row().cells
                r_row[0].merge(r_row[4])
                p = r_row[0].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(label)
                r.bold = True

                p = r_row[5].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run(f"{amount:,.2f}")
                r.bold = True

        # Remarks
        doc.add_paragraph()
        remarks = doc.add_paragraph()
        r = remarks.add_run(data.get("notes") or "")
        r.italic = True
        r.font.size = Pt(9)

        # Signature
        doc.add_paragraph()
        doc.add_paragraph()

        sig = doc.add_table(rows=1, cols=2)
        left = sig.rows[0].cells[0]
        right = sig.rows[0].cells[1]

        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Juliana\n").italic = True
        p.add_run("_________________________\n")
        p.add_run("Authorized Signature").font.size = Pt(8)

        p = right.paragraphs[0]
        p.add_run(
            "BANK INFORMATION:\n"
            "Bank Name: The Hongkong and Shanghai Banking Corporation Limited\n"
            "Bank Address: 1 Queen's Road Central, Hong Kong\n"
            "Bank Swift Code: HSBCHKHHHKH\n"
            "Account Name: Leapstack International Limited\n"
            "Account Number: 004-809-492960-838\n"
        ).font.size = Pt(8)

        # Payment terms
        doc.add_paragraph()
        terms = doc.add_paragraph()
        r = terms.add_run(data.get("payment_terms") or "")
        r.font.size = Pt(8)
        r.italic = True

        # Save
        filename = f"Invoice_{data['invoice_number']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = self.output_dir / filename
        doc.save(filepath)
        return str(filepath)