"""
Annual Report Excel generator.

- Renders one row per top-level category.
- For reimbursement categories, expands into indented sub-rows (one per
  line-item category: Transportation, Hotel, Entertainment, ...).
- Optionally pre-seeds with every category the company has ever used,
  so categories that had zero activity this year still appear.
"""
from datetime import datetime
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# ---- styles ----
TITLE_FONT = Font(bold=True, size=14)
COMPANY_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=11)
ROW_FONT = Font(size=10)
TOTAL_FONT = Font(bold=True, size=11)

HEADER_FILL = PatternFill("solid", fgColor="1A4D8F")
SECTION_FILL = PatternFill("solid", fgColor="E8F0FB")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

MONEY_FMT = '#,##0.00'


def _split_period(period: str | None) -> Tuple[str, str]:
    if not period or not isinstance(period, str):
        return ("—", "—")
    s = period.strip()
    if len(s) >= 7 and s[4] == "-":
        y, m = s[:4], s[5:7]
        if y.isdigit() and m.isdigit():
            return (y, m)
    if len(s) == 4 and s.isdigit():
        return (s, "—")
    return ("—", "—")


def _is_reimbursement(category_name: str | None, nature: str | None) -> bool:
    hay = f"{category_name or ''} {nature or ''}".lower()
    return "reimburse" in hay


def build_company_annual_sheet(
    wb,
    company_name: str,
    year: str,
    requisitions: List,
    line_items_by_req: Dict[str, List],
    known_categories: List[str] | None = None,
    known_reimb_subcategories: List[str] | None = None,
):
    sheet_name = (company_name[:28] + "..") if len(company_name) > 31 else company_name
    ws = wb.create_sheet(title=sheet_name)

    # ---------- Title block ----------
    ws.merge_cells("A1:N1")
    ws["A1"].value = company_name
    ws["A1"].font = COMPANY_FONT
    ws["A1"].alignment = LEFT

    ws.merge_cells("A2:N2")
    ws["A2"].value = f"Annual Report {year}"
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = LEFT

    # ---------- Build grid ----------
    # grid[top_cat][sub_cat][month_index] = amount
    grid: Dict[str, Dict[str, List[float]]] = {}
    unknown_month: Dict[str, Dict[str, float]] = {}

    for req in requisitions:
        top_cat = (req.category_name or req.nature or "Uncategorized").strip()
        _, month = _split_period(req.period)
        is_reimb = _is_reimbursement(req.category_name, req.nature)

        if is_reimb:
            lines = line_items_by_req.get(str(req.id), [])
            if not lines:
                sub_cat = "Miscellaneous"
                amount = float(req.total_amount or 0)
                if month.isdigit():
                    grid.setdefault(top_cat, {}).setdefault(sub_cat, [0.0]*12)[int(month)-1] += amount
                else:
                    unknown_month.setdefault(top_cat, {})[sub_cat] = (
                        unknown_month.get(top_cat, {}).get(sub_cat, 0) + amount
                    )
                continue

            for li in lines:
                sub_cat = (li.category or "Miscellaneous").strip() or "Miscellaneous"
                amount = float(li.amount_hkd or 0)
                if month.isdigit():
                    grid.setdefault(top_cat, {}).setdefault(sub_cat, [0.0]*12)[int(month)-1] += amount
                else:
                    unknown_month.setdefault(top_cat, {})[sub_cat] = (
                        unknown_month.get(top_cat, {}).get(sub_cat, 0) + amount
                    )
        else:
            sub_cat = top_cat
            amount = float(req.total_amount or 0)
            if month.isdigit():
                grid.setdefault(top_cat, {}).setdefault(sub_cat, [0.0]*12)[int(month)-1] += amount
            else:
                unknown_month.setdefault(top_cat, {})[sub_cat] = (
                    unknown_month.get(top_cat, {}).get(sub_cat, 0) + amount
                )

    # Pre-seed with known categories (even if this year had zero)
    if known_categories:
        for cat in known_categories:
            grid.setdefault(cat, {})
            if _is_reimbursement(cat, cat) and known_reimb_subcategories:
                for sub in known_reimb_subcategories:
                    grid[cat].setdefault(sub, [0.0] * 12)
            else:
                grid[cat].setdefault(cat, [0.0] * 12)

    # ---------- Header row ----------
    header_row = 4
    ws.cell(row=header_row, column=1, value="Category")
    for i, m in enumerate(MONTHS):
        ws.cell(row=header_row, column=2 + i, value=m)
    total_col = 2 + len(MONTHS)
    ws.cell(row=header_row, column=total_col, value="Total")

    for col in range(1, total_col + 1):
        c = ws.cell(row=header_row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER

    ws.row_dimensions[header_row].height = 24

    # ---------- Data rows ----------
    r = header_row + 1
    grand_monthly = [0.0] * 12

    for top_cat in sorted(grid.keys()):
        sub_cats = grid[top_cat]
        is_reimb_group = _is_reimbursement(top_cat, top_cat)

        # Section band only for reimbursement groups
        if is_reimb_group:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col)
            section = ws.cell(row=r, column=1, value=f"▸ {top_cat}")
            section.font = SECTION_FONT
            section.fill = SECTION_FILL
            section.alignment = LEFT
            for col in range(1, total_col + 1):
                ws.cell(row=r, column=col).border = BORDER
                ws.cell(row=r, column=col).fill = SECTION_FILL
            r += 1

        for sub_cat in sorted(sub_cats.keys()):
            amounts = sub_cats[sub_cat]
            label = f"    {sub_cat}" if is_reimb_group else sub_cat
            cell_label = ws.cell(row=r, column=1, value=label)
            cell_label.alignment = LEFT
            cell_label.font = ROW_FONT

            for i, amt in enumerate(amounts):
                cell = ws.cell(row=r, column=2 + i,
                               value=round(amt, 2) if amt else None)
                if amt:
                    cell.number_format = MONEY_FMT
                cell.alignment = RIGHT
                grand_monthly[i] += amt

            row_total = ws.cell(row=r, column=total_col,
                                value=f"=SUM(B{r}:M{r})")
            row_total.number_format = MONEY_FMT
            row_total.alignment = RIGHT
            row_total.font = ROW_FONT

            for col in range(1, total_col + 1):
                ws.cell(row=r, column=col).border = BORDER
            r += 1

    # ---------- Unknown-month block ----------
    if unknown_month:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=total_col)
        ws.cell(row=r, column=1, value="Items with no month").font = Font(italic=True)
        r += 1
        for top_cat, subs in sorted(unknown_month.items()):
            for sub_cat, amt in sorted(subs.items()):
                ws.cell(row=r, column=1, value=f"    {top_cat} / {sub_cat}")
                cell = ws.cell(row=r, column=total_col, value=round(amt, 2))
                cell.number_format = MONEY_FMT
                cell.alignment = RIGHT
                for col in range(1, total_col + 1):
                    ws.cell(row=r, column=col).border = BORDER
                r += 1

    # ---------- Grand total ----------
    r += 1
    ws.cell(row=r, column=1, value="TOTAL").font = TOTAL_FONT
    for i, amt in enumerate(grand_monthly):
        cell = ws.cell(row=r, column=2 + i,
                       value=round(amt, 2) if amt else None)
        if amt:
            cell.number_format = MONEY_FMT
        cell.font = TOTAL_FONT
        cell.alignment = RIGHT
    grand = ws.cell(row=r, column=total_col,
                    value=f"=SUM(B{r}:M{r})")
    grand.number_format = MONEY_FMT
    grand.font = TOTAL_FONT
    grand.alignment = RIGHT

    for col in range(1, total_col + 1):
        ws.cell(row=r, column=col).fill = TOTAL_FILL
        ws.cell(row=r, column=col).border = BORDER

    # ---------- Layout ----------
    ws.column_dimensions["A"].width = 38
    for i in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.freeze_panes = "B5"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True