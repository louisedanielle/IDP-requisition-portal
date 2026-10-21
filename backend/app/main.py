from fastapi import FastAPI, UploadFile, File, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from datetime import datetime
from uuid import UUID
from .services.revenue_service import RevenueService
from .services.invoice_generator import InvoiceGenerator
from .models import Income, IncomeLineItem
import os

from .config import settings
from .database import get_db, init_db
from .models import (
    Company, PaymentType, PaymentCategory, Payee,
    Document, ExtractedData, Requisition, LineItem,
)
from .services.mistral_service import EXPENSE_CATEGORIES
from .services.word_generator import WordGenerator
from .workers.processor import process_document


app = FastAPI(title=settings.APP_NAME)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "https://reimbursement-frontend.onrender.com",
        "https://reimbursement-frontend-*.onrender.com",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    init_db()
    print(f"✅ Application started: {settings.APP_NAME}")
    print(f"🌐 API: http://localhost:8000")


# ============ SYSTEM ============

@app.get("/")
async def root():
    return {"name": settings.APP_NAME, "status": "running"}


@app.get("/health")
async def health():
    return {"status": "healthy"}


# ============ LOOKUPS ============

@app.get("/api/categories")
async def get_expense_categories():
    return {"categories": [c["name"] for c in EXPENSE_CATEGORIES]}


@app.get("/api/companies")
async def list_companies(db: Session = Depends(get_db)):
    rows = db.query(Company).filter(Company.is_active == True).all()
    return [
        {"id": str(c.id), "name": c.name, "short_name": c.short_name}
        for c in rows
    ]


@app.get("/api/companies/{company_id}/categories")
async def list_company_categories(
    company_id: UUID, db: Session = Depends(get_db)
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    rows = (
        db.query(PaymentCategory)
        .filter(
            PaymentCategory.company_id == company_id,
            PaymentCategory.is_active == True,
        )
        .all()
    )

    grouped: dict = {}
    for c in rows:
        ptype = c.payment_type.name
        grouped.setdefault(ptype, []).append({
            "id": str(c.id),
            "name": c.name,
            "frequency": c.frequency,
        })

    order = ["Monthly Payment", "Annual Payment", "AD HOC"]
    return {k: grouped.get(k, []) for k in order if k in grouped}


@app.post("/api/companies/{company_id}/categories")
async def create_company_category(
    company_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    name = (data.get("name") or "").strip()
    payment_type_name = (data.get("payment_type") or "").strip()

    if not name:
        raise HTTPException(400, "name is required")
    if not payment_type_name:
        raise HTTPException(400, "payment_type is required")

    payment_type = (
        db.query(PaymentType)
        .filter(PaymentType.name == payment_type_name)
        .first()
    )
    if not payment_type:
        raise HTTPException(400, f"Invalid payment_type: {payment_type_name}")

    existing = (
        db.query(PaymentCategory)
        .filter(
            PaymentCategory.company_id == company.id,
            PaymentCategory.payment_type_id == payment_type.id,
            PaymentCategory.name == name,
        )
        .first()
    )
    if existing:
        return {
            "id": str(existing.id),
            "name": existing.name,
            "frequency": existing.frequency,
            "payment_type": payment_type.name,
        }

    frequency_map = {
        "Monthly Payment": "monthly",
        "Annual Payment": "annual",
        "AD HOC": "adhoc",
    }
    frequency = frequency_map.get(payment_type.name, "monthly")

    new_cat = PaymentCategory(
        company_id=company.id,
        payment_type_id=payment_type.id,
        name=name,
        frequency=frequency,
        is_active=True,
    )
    db.add(new_cat)
    db.commit()
    db.refresh(new_cat)

    print(f"➕ New Nature: {company.name} → {payment_type.name} → {name}")

    return {
        "id": str(new_cat.id),
        "name": new_cat.name,
        "frequency": new_cat.frequency,
        "payment_type": payment_type.name,
    }


@app.post("/api/companies")
async def create_company(data: dict, db: Session = Depends(get_db)):
    name = (data.get("name") or "").strip()
    short_name = (data.get("short_name") or "").strip().upper()

    if not name:
        raise HTTPException(400, "Company name is required")
    if not short_name:
        short_name = name[:3].upper()

    existing = db.query(Company).filter(Company.name == name).first()
    if existing:
        return {"id": str(existing.id), "name": existing.name, "short_name": existing.short_name}

    company = Company(name=name, short_name=short_name)
    db.add(company)
    db.commit()
    db.refresh(company)

    return {"id": str(company.id), "name": company.name, "short_name": company.short_name}


@app.get("/api/payees")
async def list_payees(db: Session = Depends(get_db)):
    rows = (
        db.query(Payee)
        .filter(Payee.is_active == True)
        .order_by(Payee.name)
        .all()
    )
    return [
        {"id": str(p.id), "name": p.name, "type": p.payee_type}
        for p in rows
    ]


@app.post("/api/payees")
async def create_payee(data: dict, db: Session = Depends(get_db)):
    name = (data.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "name is required")

    existing = db.query(Payee).filter(Payee.name == name).first()
    if existing:
        return {
            "id": str(existing.id),
            "name": existing.name,
            "type": existing.payee_type,
        }

    payee = Payee(name=name, payee_type=data.get("type", "other"))
    db.add(payee)
    db.commit()
    return {"id": str(payee.id), "name": payee.name, "type": payee.payee_type}


# ============ DOCUMENTS ============

@app.post("/api/documents/upload")
async def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    allowed = [
        "image/jpeg", "image/png", "image/tiff",
        "image/bmp", "image/webp", "application/pdf",
    ]
    if file.content_type not in allowed:
        raise HTTPException(400, f"File type not allowed. Allowed: {allowed}")

    content = await file.read()
    if not content:
        raise HTTPException(400, "Empty file")

    document = Document(
        filename=file.filename,
        file_path=None,
        file_data=content,
        file_size=len(content),
        mime_type=file.content_type,
        status="pending",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    print(f"📄 Uploaded: {document.filename} ({len(content)} bytes)")

    try:
        process_document.delay(str(document.id))
        print(f"✅ Task queued for {document.id}")
    except Exception as e:
        print(f"⚠️ Celery enqueue failed: {e}")
        document.status = "failed"
        document.error_message = f"Queue error: {e}"
        db.commit()

    return {
        "document_id": str(document.id),
        "filename": file.filename,
        "status": "processing",
    }


@app.post("/api/documents/manual")
async def create_manual_document(db: Session = Depends(get_db)):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    document = Document(
        filename=f"manual_{timestamp}",
        file_path=None,
        file_data=None,
        file_size=0,
        mime_type="application/x-manual",
        status="completed",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    extracted = ExtractedData(
        document_id=document.id,
        raw_text="",
        structured_data={
            "items": [],
            "payee_name": None,
            "document_type": "manual",
            "total_amount": None,
        },
        confidence_score=1.0,
        is_verified=False,
    )
    db.add(extracted)
    db.commit()

    print(f"✍️  Manual document created: {document.id}")

    return {
        "document_id": str(document.id),
        "status": "ready",
    }


@app.get("/api/documents")
async def list_documents(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    documents = (
        db.query(Document)
        .order_by(Document.upload_date.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [
        {
            "id": str(d.id),
            "filename": d.filename,
            "status": d.status,
            "error_message": d.error_message,
            "upload_date": d.upload_date.isoformat() if d.upload_date else None,
            "file_size": d.file_size,
            "mime_type": d.mime_type,
        }
        for d in documents
    ]


@app.get("/api/documents/{doc_id}/status")
async def get_status(doc_id: UUID, db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == doc_id).first()
    if not document:
        raise HTTPException(404, "Document not found")
    return {
        "document_id": str(document.id),
        "filename": document.filename,
        "status": document.status,
        "error_message": document.error_message,
    }


@app.get("/api/documents/{doc_id}/extracted")
async def get_extracted(doc_id: UUID, db: Session = Depends(get_db)):
    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if not extracted:
        return {
            "status": "pending",
            "items": [],
            "payee_name": None,
            "document_type": None,
            "total_amount": None,
            "raw_text": "",
        }

    structured = extracted.structured_data or {}
    return {
        "items": structured.get("items", []),
        "payee_name": structured.get("payee_name"),
        "document_type": structured.get("document_type"),
        "total_amount": structured.get("total_amount"),
        "raw_text": extracted.raw_text or "",
        "confidence": extracted.confidence_score,
        "verified": extracted.is_verified,
    }


# ============ SUGGESTIONS ============

@app.get("/api/documents/{doc_id}/suggestions")
async def get_suggestions(doc_id: UUID, db: Session = Depends(get_db)):
    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if not extracted:
        return _empty_suggestions()

    structured = extracted.structured_data or {}
    items = structured.get("items", [])
    ai_payee = (structured.get("payee_name") or "").strip() or None
    raw_text = (extracted.raw_text or "").lower()

    nature_name = _detect_nature_from_text(raw_text)
    canonical_payee = _resolve_payee(ai_payee)

    if not nature_name and canonical_payee:
        nature_name = _nature_for_payee(canonical_payee)

    if not nature_name:
        nature_name = "Reimbursement"

    matches = _find_nature_in_tree(nature_name)

    company_short = None
    payment_type_name = None
    final_payee = None

    preferred = None
    if canonical_payee:
        for m in matches:
            if m["payees"] and canonical_payee in m["payees"]:
                preferred = m
                break

    chosen = preferred or (matches[0] if matches else None)

    if chosen:
        company_short = chosen["company_short"]
        payment_type_name = chosen["payment_type"]
        payees = chosen["payees"]

        if len(payees) == 1:
            final_payee = payees[0]
        elif len(payees) > 1:
            if canonical_payee and canonical_payee in payees:
                final_payee = canonical_payee
            else:
                final_payee = payees[0]
        else:
            final_payee = canonical_payee

    if not company_short:
        company_short = "LSI"
        payment_type_name = payment_type_name or "Monthly Payment"
        final_payee = canonical_payee

    company = (
        db.query(Company).filter(Company.short_name == company_short).first()
    )

    payment_category = None
    if company and nature_name and payment_type_name:
        payment_category = (
            db.query(PaymentCategory)
            .join(PaymentType)
            .filter(
                PaymentCategory.company_id == company.id,
                PaymentCategory.name == nature_name,
                PaymentType.name == payment_type_name,
            )
            .first()
        )

    period = _infer_period(items)

    return {
        "company_id": str(company.id) if company else None,
        "company_name": company.name if company else None,
        "payment_type": payment_type_name,
        "payment_category_id": str(payment_category.id) if payment_category else None,
        "payment_category_name": payment_category.name if payment_category else None,
        "payee_name": final_payee,
        "ai_payee_name": ai_payee,
        "period": period,
        "confidence": 1.0 if (final_payee and period) else 0.5,
        "detected_nature": nature_name,
    }


# ============ REVIEW ============

@app.post("/api/documents/{doc_id}/review")
async def save_review(
    doc_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    items = data.get("items", [])

    company_id = _to_uuid(data.get("company_id"))
    payment_category_id = _to_uuid(data.get("payment_category_id"))
    payee_name = (data.get("payee_name") or "").strip()
    period = _normalize_period(data.get("period"))

    if not payee_name:
        raise HTTPException(400, "payee_name is required")
    if not company_id or not payment_category_id:
        raise HTTPException(400, "company_id and payment_category_id are required")

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    category = (
        db.query(PaymentCategory)
        .filter(PaymentCategory.id == payment_category_id)
        .first()
    )
    if not category:
        raise HTTPException(404, "Payment category not found")

    payee = db.query(Payee).filter(Payee.name == payee_name).first()
    if not payee:
        payee = Payee(name=payee_name, payee_type="other")
        db.add(payee)
        db.commit()

    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if extracted:
        existing_structured = extracted.structured_data or {}
        extracted.structured_data = {
            "items": items,
            "payee_name": payee_name,
            "company_id": str(company_id),
            "payment_category_id": str(payment_category_id),
            "period": period,
            "document_type": existing_structured.get("document_type"),
            "total_amount": existing_structured.get("total_amount"),
        }
        extracted.is_verified = True
        db.commit()

    sub_group = (data.get("sub_group") or "").strip() or None

    requisition = (
        db.query(Requisition)
        .filter(
            Requisition.company_id == company.id,
            Requisition.payment_category_id == category.id,
            Requisition.payee_name == payee_name,
            Requisition.period == period,
            Requisition.sub_group == sub_group,
        )
        .first()
    )

    if requisition is None:
        last_error = None
        for attempt in range(3):
            try:
                new_pr = _generate_pr_number(db, company.short_name)
                requisition = Requisition(
                    pr_number=new_pr,
                    company_id=company.id,
                    payment_category_id=category.id,
                    payee_id=payee.id,
                    company_name=company.name,
                    category_name=category.name,
                    payment_type_name=category.payment_type.name,
                    payee_name=payee_name,
                    period=period,
                    frequency=category.frequency,
                    sub_group=sub_group,
                    applicant="Juliana",
                    nature="Payment Requisition",
                    total_amount=0.0,
                    status="draft",
                )
                db.add(requisition)
                db.commit()
                print(
                    f"✨ Created {requisition.pr_number} "
                    f"({company.name} / {category.name} / "
                    f"{payee_name} / {period or 'no period'})"
                )
                break
            except IntegrityError as e:
                db.rollback()
                last_error = e
                print(f"⚠️ PR number collision on attempt {attempt + 1}, retrying...")
                continue

        if requisition is None or requisition.id is None:
            raise HTTPException(
                500,
                f"Could not create requisition after 3 attempts: {last_error}",
            )
    else:
        print(f"➕ Appending to {requisition.pr_number}")

    db.query(LineItem).filter(LineItem.source_document_id == doc_id).delete()
    db.flush()

    remaining = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == requisition.id)
        .order_by(LineItem.item_number)
        .all()
    )
    for i, li in enumerate(remaining, start=1):
        li.item_number = i

    real_items = [it for it in items if not it.get("is_total")]

    doc_total = None
    if extracted and extracted.structured_data:
        doc_total = extracted.structured_data.get("total_amount")

    if not real_items and doc_total is not None and float(doc_total) > 0:
        real_items = [{
            "description": _default_total_description(
                nature_name=category.name,
                doc_type=(extracted.structured_data or {}).get("document_type"),
            ),
            "amount_hkd": float(doc_total),
            "date": None,
            "category": None,
            "is_total": False,
        }]

    start_num = len(remaining) + 1
    for idx, item in enumerate(real_items):
        amount = float(item.get("amount_hkd", 0) or 0)
        line_item = LineItem(
            requisition_id=requisition.id,
            source_document_id=doc_id,
            item_number=start_num + idx,
            date=_parse_date(item.get("date")),
            description=item.get("description", ""),
            amount_hkd=amount,
            category=item.get("category") or None,
        )
        db.add(line_item)

    db.flush()

    all_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == requisition.id)
        .all()
    )
    requisition.total_amount = sum(li.amount_hkd for li in all_items)
    db.commit()

    try:
        _regenerate_word(db, requisition)
    except Exception as e:
        print(f"⚠️ Word regeneration failed: {e}")

    return {
        "message": "Review saved",
        "items_count": len(real_items),
        "requisition_id": str(requisition.id),
        "pr_number": requisition.pr_number,
    }


@app.post("/api/documents/{doc_id}/reprocess")
async def reprocess_document(doc_id: UUID, db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == doc_id).first()
    if not document:
        raise HTTPException(404, "Document not found")

    document.status = "pending"
    document.error_message = None
    db.commit()

    process_document.delay(str(document.id))
    return {"message": "Reprocessing started"}


# ============ DASHBOARD ============

@app.get("/api/dashboard/stats")
async def get_dashboard_stats(db: Session = Depends(get_db)):
    from dateutil.relativedelta import relativedelta

    today = datetime.utcnow().date()
    prev_month_date = today.replace(day=1) - relativedelta(months=1)
    prev_month_key = prev_month_date.strftime("%Y-%m")
    current_month_key = today.strftime("%Y-%m")
    current_year_key = today.strftime("%Y")

    last_rows = (
        db.query(Requisition)
        .filter(Requisition.period == prev_month_key)
        .all()
    )
    last_count = len(last_rows)
    last_total = sum(r.total_amount or 0 for r in last_rows)

    month_rows = (
        db.query(Requisition)
        .filter(Requisition.period == current_month_key)
        .all()
    )
    month_count = len(month_rows)
    month_total = sum(r.total_amount or 0 for r in month_rows)

    year_rows = (
        db.query(Requisition)
        .filter(Requisition.period.like(f"{current_year_key}%"))
        .all()
    )
    year_count = len(year_rows)
    year_total = sum(r.total_amount or 0 for r in year_rows)

    by_company = {}
    for r in year_rows:
        company = r.company_name or "Unknown"
        category = r.category_name or "Uncategorized"
        amount = r.total_amount or 0

        by_company.setdefault(company, {
            "total": 0.0,
            "count": 0,
            "categories": {},
            "company_id": str(r.company_id) if r.company_id else None,
        })
        by_company[company]["total"] += amount
        by_company[company]["count"] += 1

        by_company[company]["categories"].setdefault(category, {
            "count": 0,
            "total": 0.0,
        })
        by_company[company]["categories"][category]["count"] += 1
        by_company[company]["categories"][category]["total"] += amount

    ordered = {}
    for company in sorted(by_company.keys()):
        entry = by_company[company]
        entry["total"] = round(entry["total"], 2)

        sorted_cats = dict(
            sorted(
                entry["categories"].items(),
                key=lambda kv: kv[1]["total"],
                reverse=True,
            )
        )
        for cat in sorted_cats:
            sorted_cats[cat]["total"] = round(sorted_cats[cat]["total"], 2)

        ordered[company] = {
            "total": entry["total"],
            "count": entry["count"],
            "company_id": entry["company_id"],
            "categories": sorted_cats,
        }

    return {
        "last_month": {
            "period": prev_month_key,
            "label": prev_month_date.strftime("%B %Y"),
            "count": last_count,
            "total": round(last_total, 2),
        },
        "this_month": {
            "period": current_month_key,
            "label": today.strftime("%B %Y"),
            "count": month_count,
            "total": round(month_total, 2),
        },
        "this_year": {
            "period": current_year_key,
            "label": current_year_key,
            "count": year_count,
            "total": round(year_total, 2),
        },
        "by_company_this_year": ordered,
    }


# ============ ANNUAL REPORT EXCEL ============

def _build_annual_report_sheet(wb, company_name: str, year: str, requisitions: list):
    import openpyxl
    from openpyxl.styles import Font, Border, Side, Alignment

    sheet_name = (company_name[:28] + "..") if len(company_name) > 31 else company_name
    ws = wb.create_sheet(title=sheet_name)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    right = Alignment(horizontal="right")
    center = Alignment(horizontal="center")

    months = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    ws.merge_cells("A1:N1")
    ws["A1"] = f"{company_name} — Annual Report {year}"
    ws["A1"].font = Font(bold=True, size=14)

    header_row = 3
    ws.cell(row=header_row, column=1, value="Category").font = bold
    for i, m in enumerate(months):
        c = ws.cell(row=header_row, column=2 + i, value=m)
        c.font = bold
        c.alignment = center
    total_col = 2 + len(months)
    ws.cell(row=header_row, column=total_col, value="Total").font = bold
    ws.cell(row=header_row, column=total_col).alignment = center

    grid: dict = {}
    unknown_month: dict = {}

    for r in requisitions:
        category = r.category_name or "Uncategorized"
        amount = float(r.total_amount or 0)

        _, month = _split_period(r.period)
        if month == "—" or not month.isdigit():
            unknown_month[category] = unknown_month.get(category, 0) + amount
            continue

        m_idx = int(month) - 1
        if m_idx < 0 or m_idx > 11:
            unknown_month[category] = unknown_month.get(category, 0) + amount
            continue

        grid.setdefault(category, [0.0] * 12)
        grid[category][m_idx] += amount

    row = header_row + 1
    for category in sorted(grid.keys()):
        ws.cell(row=row, column=1, value=category)
        amounts = grid[category]
        for i, amt in enumerate(amounts):
            cell = ws.cell(row=row, column=2 + i, value=round(amt, 2))
            cell.number_format = "#,##0.00"
            cell.alignment = right
        total_cell = ws.cell(
            row=row,
            column=total_col,
            value=f"=SUM(B{row}:M{row})",
        )
        total_cell.font = bold
        total_cell.number_format = "#,##0.00"
        total_cell.alignment = right
        row += 1

    if unknown_month:
        row += 1
        ws.cell(row=row, column=1, value="Items with no month").font = Font(italic=True)
        row += 1
        for category, amt in sorted(unknown_month.items()):
            ws.cell(row=row, column=1, value=category)
            ws.cell(row=row, column=total_col, value=round(amt, 2)).number_format = "#,##0.00"
            row += 1

    total_row = row + 1
    ws.cell(row=total_row, column=1, value="Total").font = bold
    for i in range(2, 2 + len(months)):
        col_letter = openpyxl.utils.get_column_letter(i)
        c = ws.cell(
            row=total_row,
            column=i,
            value=f"=SUM({col_letter}{header_row + 1}:{col_letter}{row})",
        )
        c.font = bold
        c.number_format = "#,##0.00"
        c.alignment = right
    grand_total = ws.cell(
        row=total_row,
        column=total_col,
        value=f"=SUM(N{header_row + 1}:N{row})",
    )
    grand_total.font = bold
    grand_total.number_format = "#,##0.00"
    grand_total.alignment = right

    ws.column_dimensions["A"].width = 32
    for i in range(2, 2 + len(months) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12

    for r in range(header_row, total_row + 1):
        for c in range(1, total_col + 1):
            ws.cell(row=r, column=c).border = border


@app.get("/api/companies/{company_id}/annual-report")
async def company_annual_report(
    company_id: UUID,
    year: str | None = None,
    db: Session = Depends(get_db),
):
    import openpyxl
    import tempfile

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    year = (year or datetime.utcnow().strftime("%Y")).strip()

    rows = (
        db.query(Requisition)
        .filter(
            Requisition.company_id == company.id,
            Requisition.period.like(f"{year}%"),
        )
        .order_by(Requisition.created_date.asc())
        .all()
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_annual_report_sheet(wb, company.name, year, rows)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    safe_name = "".join(
        c if c.isalnum() or c in " -_" else "_"
        for c in company.name
    )[:60]

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{safe_name} - Annual Report {year}.xlsx",
    )


@app.get("/api/dashboard/annual-report")
async def dashboard_annual_report(
    year: str | None = None,
    db: Session = Depends(get_db),
):
    import openpyxl
    import tempfile
    from openpyxl.styles import Font, Border, Side, Alignment

    year = (year or datetime.utcnow().strftime("%Y")).strip()

    companies = (
        db.query(Company)
        .filter(Company.is_active == True)
        .order_by(Company.name)
        .all()
    )

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = f"Summary {year}"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    right = Alignment(horizontal="right")
    center = Alignment(horizontal="center")

    months = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    summary.merge_cells("A1:N1")
    summary["A1"] = f"All Companies — Annual Report {year}"
    summary["A1"].font = Font(bold=True, size=14)

    header_row = 3
    summary.cell(row=header_row, column=1, value="Company").font = bold
    for i, m in enumerate(months):
        c = summary.cell(row=header_row, column=2 + i, value=m)
        c.font = bold
        c.alignment = center
    total_col = 2 + len(months)
    summary.cell(row=header_row, column=total_col, value="Total").font = bold
    summary.cell(row=header_row, column=total_col).alignment = center

    row = header_row + 1
    for company in companies:
        rows = (
            db.query(Requisition)
            .filter(
                Requisition.company_id == company.id,
                Requisition.period.like(f"{year}%"),
            )
            .all()
        )

        if rows:
            _build_annual_report_sheet(wb, company.name, year, rows)

        amounts = [0.0] * 12
        for r in rows:
            _, month = _split_period(r.period)
            if not month.isdigit():
                continue
            m_idx = int(month) - 1
            if 0 <= m_idx < 12:
                amounts[m_idx] += float(r.total_amount or 0)

        summary.cell(row=row, column=1, value=company.name)
        for i, amt in enumerate(amounts):
            cell = summary.cell(row=row, column=2 + i, value=round(amt, 2))
            cell.number_format = "#,##0.00"
            cell.alignment = right
        total_cell = summary.cell(
            row=row, column=total_col, value=f"=SUM(B{row}:M{row})"
        )
        total_cell.font = bold
        total_cell.number_format = "#,##0.00"
        total_cell.alignment = right
        row += 1

    total_row = row
    summary.cell(row=total_row, column=1, value="Total").font = bold
    for i in range(2, 2 + len(months)):
        col = openpyxl.utils.get_column_letter(i)
        c = summary.cell(
            row=total_row,
            column=i,
            value=f"=SUM({col}{header_row + 1}:{col}{row - 1})",
        )
        c.font = bold
        c.number_format = "#,##0.00"
        c.alignment = right
    gt = summary.cell(
        row=total_row, column=total_col, value=f"=SUM(N{header_row + 1}:N{row - 1})"
    )
    gt.font = bold
    gt.number_format = "#,##0.00"
    gt.alignment = right

    summary.column_dimensions["A"].width = 36
    for i in range(2, 2 + len(months) + 1):
        summary.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12

    for r in range(header_row, total_row + 1):
        for c in range(1, total_col + 1):
            summary.cell(row=r, column=c).border = border

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Annual Report {year}.xlsx",
    )


# ============ REQUISITIONS ============

@app.get("/api/requisitions")
async def list_requisitions(
    company_id: UUID | None = None,
    category_id: UUID | None = None,
    payee_name: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(Requisition)
    if company_id:
        q = q.filter(Requisition.company_id == company_id)
    if category_id:
        q = q.filter(Requisition.payment_category_id == category_id)
    if payee_name:
        q = q.filter(Requisition.payee_name == payee_name)
    if status:
        q = q.filter(Requisition.status == status)

    rows = q.order_by(Requisition.created_date.desc()).all()

    result = []
    for r in rows:
        item_count = (
            db.query(LineItem)
            .filter(LineItem.requisition_id == r.id)
            .count()
        )
        result.append({
            "id": str(r.id),
            "pr_number": r.pr_number,
            "company_name": r.company_name,
            "category_name": r.category_name,
            "payment_type_name": r.payment_type_name,
            "payee_name": r.payee_name,
            "period": r.period,
            "total_amount": r.total_amount,
            "status": r.status,
            "item_count": item_count,
            "created_date": r.created_date.isoformat() if r.created_date else None,
        })
    return result


@app.get("/api/requisitions/folders")
async def list_folder_view(
    company_id: UUID | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(Requisition)
    if company_id:
        q = q.filter(Requisition.company_id == company_id)

    rows = q.order_by(Requisition.created_date.desc()).all()

    tree: dict = {}
    for r in rows:
        company = r.company_name or "Unknown Company"
        year, month = _split_period(r.period)
        category = r.category_name or "Uncategorized"
        payee = r.payee_name or "Unknown Payee"

        tree.setdefault(company, {})
        tree[company].setdefault(year, {})
        tree[company][year].setdefault(month, {})
        tree[company][year][month].setdefault(category, {})
        tree[company][year][month][category].setdefault(payee, []).append({
            "id": str(r.id),
            "pr_number": r.pr_number,
            "payee_name": r.payee_name,
            "payment_type_name": r.payment_type_name,
            "total_amount": r.total_amount,
            "status": r.status,
            "period": r.period,
            "created_date": r.created_date.isoformat() if r.created_date else None,
            "item_count": (
                db.query(LineItem)
                .filter(LineItem.requisition_id == r.id)
                .count()
            ),
        })

    sorted_tree = {}
    for company in sorted(tree.keys()):
        sorted_tree[company] = {}
        for year in sorted(tree[company].keys(), reverse=True):
            sorted_tree[company][year] = {}
            for month in sorted(tree[company][year].keys(), reverse=True):
                sorted_tree[company][year][month] = {}
                for category in sorted(tree[company][year][month].keys()):
                    sorted_tree[company][year][month][category] = {}
                    for payee in sorted(tree[company][year][month][category].keys()):
                        prs = sorted(
                            tree[company][year][month][category][payee],
                            key=lambda x: x.get("created_date") or "",
                            reverse=True,
                        )
                        sorted_tree[company][year][month][category][payee] = prs

    return sorted_tree


@app.get("/api/requisitions/{req_id}")
async def get_requisition(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req_id)
        .order_by(LineItem.item_number)
        .all()
    )

    return {
        "id": str(req.id),
        "pr_number": req.pr_number,
        "company_name": req.company_name,
        "category_name": req.category_name,
        "payment_type_name": req.payment_type_name,
        "payee_name": req.payee_name,
        "applicant": req.applicant,
        "nature": req.nature,
        "period": req.period,
        "frequency": req.frequency,
        "sub_group": req.sub_group,
        "total_amount": req.total_amount,
        "status": req.status,
        "created_date": req.created_date.isoformat() if req.created_date else None,
        "line_items": [
            {
                "id": str(li.id),
                "item_number": li.item_number,
                "date": li.date.isoformat() if li.date else None,
                "description": li.description,
                "amount_hkd": li.amount_hkd,
                "category": li.category,
            }
            for li in line_items
        ],
    }


@app.put("/api/requisitions/{req_id}")
async def update_requisition(
    req_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    if "pr_number" in data:
        new_pr = (data.get("pr_number") or "").strip()
        if not new_pr:
            raise HTTPException(400, "pr_number cannot be empty")
        if new_pr != req.pr_number:
            clash = (
                db.query(Requisition)
                .filter(Requisition.pr_number == new_pr, Requisition.id != req.id)
                .first()
            )
            if clash:
                raise HTTPException(400, f"PR number '{new_pr}' already exists")
            req.pr_number = new_pr

    new_company_id = _to_uuid(data.get("company_id"))
    new_category_id = _to_uuid(data.get("payment_category_id"))

    if new_company_id:
        new_company = db.query(Company).filter(Company.id == new_company_id).first()
        if not new_company:
            raise HTTPException(404, "Company not found")
        req.company_id = new_company.id
        req.company_name = new_company.name

    if new_category_id:
        new_category = (
            db.query(PaymentCategory)
            .filter(PaymentCategory.id == new_category_id)
            .first()
        )
        if not new_category:
            raise HTTPException(404, "Payment category not found")
        req.payment_category_id = new_category.id
        req.category_name = new_category.name
        req.payment_type_name = new_category.payment_type.name
        req.frequency = new_category.frequency

    if "payee_name" in data:
        new_payee_name = (data.get("payee_name") or "").strip()
        if not new_payee_name:
            raise HTTPException(400, "payee_name cannot be empty")
        payee = db.query(Payee).filter(Payee.name == new_payee_name).first()
        if not payee:
            payee = Payee(name=new_payee_name, payee_type="other")
            db.add(payee)
            db.flush()
        req.payee_id = payee.id
        req.payee_name = new_payee_name

    if "period" in data:
        req.period = _normalize_period(data.get("period"))

    if "line_items" in data and data["line_items"] is not None:
        line_items = data["line_items"]

        db.query(LineItem).filter(LineItem.requisition_id == req.id).delete()
        db.flush()

        for idx, item in enumerate(line_items):
            amount = float(item.get("amount_hkd", 0) or 0)
            line_item = LineItem(
                requisition_id=req.id,
                source_document_id=item.get("source_document_id") or None,
                item_number=idx + 1,
                date=_parse_date(item.get("date")),
                description=item.get("description", "") or "",
                amount_hkd=amount,
                category=item.get("category") or None,
            )
            db.add(line_item)

        db.flush()

        all_items = (
            db.query(LineItem)
            .filter(LineItem.requisition_id == req.id)
            .all()
        )
        req.total_amount = sum(li.amount_hkd for li in all_items)

    db.commit()
    db.refresh(req)

    try:
        _regenerate_word(db, req)
    except Exception as e:
        print(f"⚠️ Word regeneration failed: {e}")

    return {
        "message": "Requisition updated",
        "id": str(req.id),
        "pr_number": req.pr_number,
        "total_amount": req.total_amount,
    }


@app.delete("/api/requisitions/{req_id}")
async def delete_requisition(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    db.query(LineItem).filter(LineItem.requisition_id == req_id).delete()
    db.delete(req)
    db.commit()
    return {"message": "Deleted successfully"}


@app.get("/api/requisitions/{req_id}/export/word")
async def export_word(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Not found")

    docx_path = _regenerate_word(db, req)

    return FileResponse(
        docx_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{req.pr_number}.docx",
    )


@app.get("/api/requisitions/{req_id}/export/excel")
async def export_excel(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Not found")

    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req_id)
        .order_by(LineItem.item_number)
        .all()
    )

    import openpyxl
    from openpyxl.styles import Font, Border, Side, Alignment
    import tempfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{req.pr_number}_Detail List"

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    period_label = req.period or ""
    payee_label = req.payee_name or ""
    title = f"Mr {payee_label}'s Reimbursement - {period_label}"
    ws.merge_cells("A1:K1")
    ws["A1"] = title
    ws["A1"].font = bold

    headers = ["", "Item", "No.", "Date", "Description", "RMB", "HKD", "RMB", "HKD", "HKD"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = bold

    groups: dict = {}
    for li in line_items:
        cat = (li.category or "Miscellaneous").strip() or "Miscellaneous"
        groups.setdefault(cat, []).append(li)

    row = 4
    for category, items in groups.items():
        cat_start_row = row
        ws.cell(row=row, column=2, value=category)

        for idx, li in enumerate(items):
            ws.cell(row=row, column=3, value=idx + 1)
            ws.cell(row=row, column=4, value=li.date if li.date else None)
            if li.date:
                ws.cell(row=row, column=4).number_format = "yyyy-mm-dd"
            desc_cell = ws.cell(row=row, column=5, value=li.description)
            desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=7, value=li.amount_hkd)
            row += 1

        cat_end_row = row - 1
        ws.cell(
            row=cat_start_row,
            column=10,
            value=f"=SUM(G{cat_start_row}:G{cat_end_row})",
        )

    grand_total_row = row + 1
    ws.cell(row=grand_total_row, column=6, value=f"=SUM(G4:G{row-1})")
    ws.cell(row=grand_total_row, column=10, value=f"=SUM(J4:J{row-1})")

    subtotal_row = grand_total_row + 2
    ws.cell(row=subtotal_row, column=5, value="HKD sub-total")
    ws.cell(row=subtotal_row, column=7, value=f"=SUM(G4:G{row-1})")

    total_row = subtotal_row + 2
    ws.cell(row=total_row, column=5, value="Total")
    ws.cell(row=total_row, column=6, value=f"=F{subtotal_row}+G{subtotal_row}")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 20

    max_col = 10
    max_row = total_row
    for r in range(3, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None or (4 <= r <= row - 1):
                cell.border = border

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{req.pr_number}_Detail List.xlsx",
    )


# ============ INCOME / INVOICES ============

revenue_service = RevenueService()
invoice_generator = InvoiceGenerator()


def _resolve_income_relations(db, company_id, payment_category_id, payee_name):
    company = db.query(Company).filter(Company.id == company_id).first() if company_id else None
    category = None
    if payment_category_id:
        category = (
            db.query(PaymentCategory)
            .filter(PaymentCategory.id == payment_category_id)
            .first()
        )
    payee = None
    if payee_name:
        name = payee_name.strip()
        payee = db.query(Payee).filter(Payee.name == name).first()
        if not payee:
            payee = Payee(name=name, payee_type="other")
            db.add(payee)
            db.flush()
    return {"company": company, "category": category, "payee": payee}


@app.post("/api/incomes/upload")
async def upload_income_report(file: UploadFile = File(...), db: Session = Depends(get_db)):
    allowed = ["image/jpeg", "image/png", "application/pdf"]
    if file.content_type not in allowed:
        raise HTTPException(400, f"File type not allowed. Allowed: {allowed}")

    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = os.path.join(settings.UPLOAD_DIR, f"{timestamp}_{file.filename}")
    content = await file.read()
    with open(filepath, "wb") as f:
        f.write(content)

    result = await revenue_service.extract_revenue_report(filepath)
    if result.get("error"):
        raise HTTPException(500, f"Extraction failed: {result['error']}")
    return result


@app.post("/api/incomes")
@app.post("/api/incomes")
async def create_income(data: dict, db: Session = Depends(get_db)):
    company_id = _to_uuid(data.get("company_id"))
    if not company_id:
        raise HTTPException(400, "company_id is required")

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(400, "Valid company_id is required")

    customer_name = (data.get("customer_name") or "").strip()
    if not customer_name:
        raise HTTPException(400, "customer_name is required")

    customer_address = (data.get("customer_address") or "").strip() or None

    requested = (data.get("invoice_number") or "").strip()
    if requested:
        invoice_number = requested
    else:
        import re as _re
        prefix = (company.short_name or "LSI").upper()
        last = (
            db.query(Income)
            .filter(Income.invoice_number.like(f"{prefix}%"))
            .order_by(Income.invoice_number.desc())
            .first()
        )
        next_num = 1
        if last:
            m = _re.search(r"(\d+)$", last.invoice_number or "")
            if m:
                next_num = int(m.group(1)) + 1
        invoice_number = f"{prefix}{next_num:04d}"

    period = _normalize_period(data.get("period"))
    sub_group = (data.get("sub_group") or "").strip() or None

    income = Income(
        invoice_number=invoice_number,
        company_id=company.id,
        company_name=company.name,
        customer_name=customer_name,
        customer_address=customer_address,
        period=period,
        date_issued=_parse_date(data.get("date_issued")) or datetime.utcnow(),
        total_rm=float(data.get("total_rm") or 0),
        total_hkd=float(data["total_hkd"]) if data.get("total_hkd") else None,
        total_usd=float(data["total_usd"]) if data.get("total_usd") else None,
        currency="RM",
        status="draft",
        statement_ref=data.get("statement_ref"),
        notes=data.get("notes"),
        sub_group=sub_group,
    )
    db.add(income)
    db.flush()

    for idx, li in enumerate(data.get("lines", data.get("line_items", [])), start=1):
        line = IncomeLineItem(
            income_id=income.id,
            item_number=idx,
            description=li.get("branch") or li.get("description") or "",
            invoice_no=li.get("invoice_no"),
            machine_serial=li.get("machine_serial"),
            machine_runs=int(li.get("machine_runs") or 0),
            rate=float(li.get("rate_rm") or li.get("rate") or 0),
            amount=float(li.get("amount_rm") or li.get("amount") or 0),
            currency="RM",
        )
        db.add(line)

    db.commit()
    return {"id": str(income.id), "invoice_number": income.invoice_number}


@app.get("/api/incomes")
async def list_incomes(
    company_id: UUID | None = None,
    payment_category_id: UUID | None = None,
    payee_name: str | None = None,
    period: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(Income)
    if company_id:
        q = q.filter(Income.company_id == company_id)
    if payment_category_id:
        q = q.filter(Income.payment_category_id == payment_category_id)
    if payee_name:
        q = q.filter(Income.payee_name == payee_name)
    if period:
        q = q.filter(Income.period == period)
    if status:
        q = q.filter(Income.status == status)

    rows = q.order_by(Income.date_issued.desc()).all()
    return [
        {
            "id": str(r.id),
            "invoice_number": r.invoice_number,
            "company_id": str(r.company_id) if r.company_id else None,
            "company_name": r.company_name,
            "customer_name": r.customer_name,
            "payee_name": r.payee_name,
            "category_name": r.category_name,
            "payment_type_name": r.payment_type_name,
            "payment_category_id": (
                str(r.payment_category_id) if r.payment_category_id else None
            ),
            "period": r.period,
            "date_issued": r.date_issued.isoformat() if r.date_issued else None,
            "currency": r.currency,
            "total_rm": r.total_rm,
            "total_hkd": r.total_hkd,
            "total_usd": r.total_usd,
            "status": r.status,
            "sub_group": r.sub_group,
        }
        for r in rows
    ]


@app.get("/api/incomes/folders")
async def list_income_folders(company_id: UUID | None = None, db: Session = Depends(get_db)):
    q = db.query(Income)
    if company_id:
        q = q.filter(Income.company_id == company_id)
    rows = q.order_by(Income.date_issued.desc()).all()

    tree = {}
    for r in rows:
        company = r.company_name or "Unknown Company"
        year, month = _split_period(r.period)
        customer = r.customer_name or "Unknown Customer"

        tree.setdefault(company, {})
        tree[company].setdefault(year, {})
        tree[company][year].setdefault(month, {})
        tree[company][year][month].setdefault(customer, []).append({
            "id": str(r.id),
            "invoice_number": r.invoice_number,
            "customer_name": r.customer_name,
            "customer_address": r.customer_address,
            "period": r.period,
            "total_rm": r.total_rm,
            "total_hkd": r.total_hkd,
            "status": r.status,
            "date_issued": r.date_issued.isoformat() if r.date_issued else None,
        })

    sorted_tree = {}
    for company in sorted(tree.keys()):
        sorted_tree[company] = {}
        for year in sorted(tree[company].keys(), reverse=True):
            sorted_tree[company][year] = {}
            for month in sorted(tree[company][year].keys(), reverse=True):
                sorted_tree[company][year][month] = {}
                for customer in sorted(tree[company][year][month].keys()):
                    invs = sorted(
                        tree[company][year][month][customer],
                        key=lambda x: x.get("date_issued") or "",
                        reverse=True,
                    )
                    sorted_tree[company][year][month][customer] = invs
    return sorted_tree


@app.get("/api/incomes/{income_id}")
async def get_income(income_id: UUID, db: Session = Depends(get_db)):
    r = db.query(Income).filter(Income.id == income_id).first()
    if not r:
        raise HTTPException(404, "Income not found")

    lines = (
        db.query(IncomeLineItem)
        .filter(IncomeLineItem.income_id == income_id)
        .order_by(IncomeLineItem.item_number)
        .all()
    )

    normalised = [
        {
            "id": str(li.id),
            "item_number": li.item_number,
            "description": li.description,
            "branch": li.description,
            "invoice_no": li.invoice_no,
            "machine_serial": li.machine_serial,
            "machine_runs": li.machine_runs,
            "rate": li.rate,
            "rate_rm": li.rate,
            "amount": li.amount,
            "amount_rm": li.amount,
        }
        for li in lines
    ]

    return {
        "id": str(r.id),
        "invoice_number": r.invoice_number,
        "company_id": str(r.company_id) if r.company_id else None,
        "company_name": r.company_name,
        "customer_name": r.customer_name,
        "customer_address": r.customer_address,
        "payee_id": str(r.payee_id) if r.payee_id else None,
        "payee_name": r.payee_name,
        "category_name": r.category_name,
        "payment_category_id": str(r.payment_category_id) if r.payment_category_id else None,
        "payment_type_name": r.payment_type_name,
        "period": r.period,
        "date_issued": r.date_issued.isoformat() if r.date_issued else None,
        "total_rm": r.total_rm,
        "total_hkd": r.total_hkd,
        "status": r.status,
        "statement_ref": r.statement_ref,
        "notes": r.notes,
        "sub_group": r.sub_group,
        "line_items": normalised,
        "lines": normalised,
    }


@app.put("/api/incomes/{income_id}")
async def update_income(income_id: UUID, data: dict, db: Session = Depends(get_db)):
    r = db.query(Income).filter(Income.id == income_id).first()
    if not r:
        raise HTTPException(404, "Income not found")

    if "invoice_number" in data:
        r.invoice_number = data["invoice_number"]

    if "customer_name" in data:
        new_cn = (data.get("customer_name") or "").strip()
        if not new_cn:
            raise HTTPException(400, "customer_name cannot be empty")
        r.customer_name = new_cn
    if "customer_address" in data:
        r.customer_address = (data.get("customer_address") or "").strip() or None

    if "period" in data:
        r.period = _normalize_period(data["period"])
    if "date_issued" in data:
        r.date_issued = _parse_date(data["date_issued"]) or datetime.utcnow()
    if "total_rm" in data:
        r.total_rm = float(data["total_rm"] or 0)
    if "total_hkd" in data:
        r.total_hkd = float(data["total_hkd"]) if data["total_hkd"] else None
    if "statement_ref" in data:
        r.statement_ref = data["statement_ref"]
    if "notes" in data:
        r.notes = data["notes"]
    if "sub_group" in data:
        r.sub_group = (data.get("sub_group") or "").strip() or None

    lines_payload = data.get("line_items", data.get("lines"))
    if lines_payload is not None:
        db.query(IncomeLineItem).filter(IncomeLineItem.income_id == r.id).delete()
        for idx, li in enumerate(lines_payload, start=1):
            line = IncomeLineItem(
                income_id=r.id,
                item_number=idx,
                description=li.get("branch") or li.get("description") or "",
                invoice_no=li.get("invoice_no"),
                machine_serial=li.get("machine_serial"),
                machine_runs=int(li.get("machine_runs") or 0),
                rate=float(li.get("rate_rm") or li.get("rate") or 0),
                amount=float(li.get("amount_rm") or li.get("amount") or 0),
                currency="RM",
            )
            db.add(line)

        if "total_rm" not in data:
            db.flush()
            all_lines = (
                db.query(IncomeLineItem)
                .filter(IncomeLineItem.income_id == r.id)
                .all()
            )
            r.total_rm = sum(li.amount for li in all_lines)

    db.commit()
    return {"message": "Updated successfully"}


@app.delete("/api/incomes/{income_id}")
async def delete_income(income_id: UUID, db: Session = Depends(get_db)):
    r = db.query(Income).filter(Income.id == income_id).first()
    if not r:
        raise HTTPException(404, "Income not found")
    db.query(IncomeLineItem).filter(IncomeLineItem.income_id == income_id).delete()
    db.delete(r)
    db.commit()
    return {"message": "Deleted successfully"}


@app.get("/api/incomes/{income_id}/export/word")
async def export_invoice(income_id: UUID, db: Session = Depends(get_db)):
    r = db.query(Income).filter(Income.id == income_id).first()
    if not r:
        raise HTTPException(404, "Income not found")

    lines = (
        db.query(IncomeLineItem)
        .filter(IncomeLineItem.income_id == income_id)
        .order_by(IncomeLineItem.item_number)
        .all()
    )

    doc_data = {
        "invoice_number": r.invoice_number,
        "company_name": r.company_name,
        "customer_name": r.customer_name,
        "customer_address": r.customer_address,
        "date_issued": r.date_issued.strftime("%d-%b-%Y") if r.date_issued else "",
        "line_items": [
            {
                "description": li.description,
                "invoice_no": li.invoice_no,
                "machine_serial": li.machine_serial,
                "machine_runs": li.machine_runs,
                "rate": li.rate,
                "amount": li.amount,
            }
            for li in lines
        ],
        "total_rm": r.total_rm,
        "total_hkd": r.total_hkd,
        "total_usd": r.total_usd,
        "notes": r.notes or "",
        "period": r.period,
        "statement_ref": r.statement_ref,
        "payment_terms": f"Monthly Revenue Sharing - {r.period or ''}. Payment due within 15 days.",
    }

    path = invoice_generator.generate_invoice_docx(doc_data)
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{r.invoice_number}.docx",
    )


@app.get("/api/dashboard/income-stats")
async def get_income_stats(year: str | None = None, db: Session = Depends(get_db)):
    year = (year or datetime.utcnow().strftime("%Y")).strip()
    rows = db.query(Income).filter(Income.period.like(f"{year}%")).all()

    months = ["01","02","03","04","05","06","07","08","09","10","11","12"]
    by_company: dict = {}
    total_rm = 0.0
    total_hkd = 0.0
    monthly_totals = {m: {"rm": 0.0, "hkd": 0.0} for m in months}

    for r in rows:
        company = r.company_name or "Unknown"
        _, month = _split_period(r.period)
        by_company.setdefault(company, {
            "total_rm": 0.0, "total_hkd": 0.0, "invoices": 0,
            "monthly": {m: {"rm": 0.0, "hkd": 0.0} for m in months},
        })
        by_company[company]["total_rm"] += r.total_rm or 0
        by_company[company]["total_hkd"] += r.total_hkd or 0
        by_company[company]["invoices"] += 1
        if month in months:
            by_company[company]["monthly"][month]["rm"] += r.total_rm or 0
            by_company[company]["monthly"][month]["hkd"] += r.total_hkd or 0
            monthly_totals[month]["rm"] += r.total_rm or 0
            monthly_totals[month]["hkd"] += r.total_hkd or 0
        total_rm += r.total_rm or 0
        total_hkd += r.total_hkd or 0

    return {
        "year": year,
        "total_rm": round(total_rm, 2),
        "total_hkd": round(total_hkd, 2),
        "by_company": by_company,
        "monthly_totals": monthly_totals,
    }


@app.get("/api/dashboard/monthly-report")
async def monthly_report_excel(year: str | None = None, db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
    import tempfile

    year = (year or datetime.utcnow().strftime("%Y")).strip()

    incomes = db.query(Income).filter(Income.period.like(f"{year}%")).all()
    expenses = (
        db.query(Requisition)
        .filter(Requisition.period.like(f"{year}%"))
        .all()
    )

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    right = Alignment(horizontal="right")
    center = Alignment(horizontal="center")

    months = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

    ws = wb.active
    ws.title = f"Summary {year}"

    ws.merge_cells("A1:O1")
    ws["A1"] = f"Income vs Expenses — {year}"
    ws["A1"].font = Font(bold=True, size=14)

    header_row = 3
    ws.cell(row=header_row, column=1, value="Type").font = bold
    ws.cell(row=header_row, column=2, value="Name").font = bold
    for i, m in enumerate(months):
        c = ws.cell(row=header_row, column=3 + i, value=m)
        c.font = bold
        c.alignment = center
    ws.cell(row=header_row, column=3 + len(months), value="Total").font = bold

    total_col = 3 + len(months)

    def add_row(r, type_label, name, monthly_values):
        ws.cell(row=r, column=1, value=type_label)
        ws.cell(row=r, column=2, value=name)
        for i, v in enumerate(monthly_values):
            c = ws.cell(row=r, column=3 + i, value=round(v, 2))
            c.number_format = "#,##0.00"
            c.alignment = right
        total = ws.cell(row=r, column=total_col, value=f"=SUM(C{r}:N{r})")
        total.font = bold
        total.number_format = "#,##0.00"
        total.alignment = right
        return r + 1

    income_by_company: dict = {}
    for r in incomes:
        company = r.company_name or "Unknown"
        _, month = _split_period(r.period)
        if month not in [f"{i+1:02d}" for i in range(12)]:
            continue
        idx = int(month) - 1
        income_by_company.setdefault(company, [0.0]*12)
        income_by_company[company][idx] += r.total_rm or 0

    r = header_row + 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="INCOME (RM)").font = Font(bold=True, size=11)
    r += 1
    for company in sorted(income_by_company.keys()):
        r = add_row(r, "Income", company, income_by_company[company])

    expense_by_cat: dict = {}
    for e in expenses:
        cat = e.category_name or "Uncategorized"
        _, month = _split_period(e.period)
        if month not in [f"{i+1:02d}" for i in range(12)]:
            continue
        idx = int(month) - 1
        expense_by_cat.setdefault(cat, [0.0]*12)
        expense_by_cat[cat][idx] += e.total_amount or 0

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.cell(row=r, column=1, value="EXPENSES (HKD)").font = Font(bold=True, size=11)
    r += 1
    for cat in sorted(expense_by_cat.keys()):
        r = add_row(r, "Expense", cat, expense_by_cat[cat])

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    for i in range(12):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3+i)].width = 12
    ws.column_dimensions[openpyxl.utils.get_column_letter(total_col)].width = 14

    for rr in range(header_row, r):
        for cc in range(1, total_col + 1):
            ws.cell(row=rr, column=cc).border = border

    wsi = wb.create_sheet("Income Detail")
    hdr = ["Invoice No", "Company", "Customer", "Period", "Date", "Currency", "Total RM", "Total HKD", "Status"]
    for i, h in enumerate(hdr, start=1):
        wsi.cell(row=1, column=i, value=h).font = bold

    for idx, inc in enumerate(incomes, start=2):
        wsi.cell(row=idx, column=1, value=inc.invoice_number)
        wsi.cell(row=idx, column=2, value=inc.company_name)
        wsi.cell(row=idx, column=3, value=inc.customer_name)
        wsi.cell(row=idx, column=4, value=inc.period)
        wsi.cell(row=idx, column=5, value=inc.date_issued.strftime("%Y-%m-%d") if inc.date_issued else "")
        wsi.cell(row=idx, column=6, value=inc.currency)
        wsi.cell(row=idx, column=7, value=inc.total_rm).number_format = "#,##0.00"
        wsi.cell(row=idx, column=8, value=inc.total_hkd or 0).number_format = "#,##0.00"
        wsi.cell(row=idx, column=9, value=inc.status)

    wsi.column_dimensions["A"].width = 18
    wsi.column_dimensions["B"].width = 30
    wsi.column_dimensions["C"].width = 30
    for col in ["D","E","F","G","H","I"]:
        wsi.column_dimensions[col].width = 14

    wse = wb.create_sheet("Expense Detail")
    hdr = ["PR No", "Company", "Category", "Payee", "Period", "Total HKD", "Status"]
    for i, h in enumerate(hdr, start=1):
        wse.cell(row=1, column=i, value=h).font = bold

    for idx, exp in enumerate(expenses, start=2):
        wse.cell(row=idx, column=1, value=exp.pr_number)
        wse.cell(row=idx, column=2, value=exp.company_name)
        wse.cell(row=idx, column=3, value=exp.category_name)
        wse.cell(row=idx, column=4, value=exp.payee_name)
        wse.cell(row=idx, column=5, value=exp.period)
        wse.cell(row=idx, column=6, value=exp.total_amount).number_format = "#,##0.00"
        wse.cell(row=idx, column=7, value=exp.status)

    wse.column_dimensions["A"].width = 18
    wse.column_dimensions["B"].width = 30
    wse.column_dimensions["C"].width = 22
    wse.column_dimensions["D"].width = 25
    wse.column_dimensions["E"].width = 12
    wse.column_dimensions["F"].width = 14
    wse.column_dimensions["G"].width = 12

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Monthly Report {year}.xlsx",
    )


# ============ HELPERS ============

def _empty_suggestions():
    return {
        "company_id": None, "payment_type": None,
        "payment_category_id": None, "payee_name": None, "period": None,
        "confidence": 0.0,
    }


def _normalize_name(s: str) -> str:
    import re as _re
    s = (s or "").lower()
    s = _re.sub(r"\b(mr|mrs|ms|dr|prof|professor)\.?\b", " ", s)
    s = _re.sub(r"[.,\-_/\\()\[\]]", " ", s)
    return " ".join(s.split())


def _resolve_payee(ai_payee: str | None) -> str | None:
    if not ai_payee:
        return None
    needle = _normalize_name(ai_payee)
    for alias, canonical in settings.PAYEE_ALIASES.items():
        if _normalize_name(alias) == needle:
            return canonical
    for alias, canonical in settings.PAYEE_ALIASES.items():
        a = _normalize_name(alias)
        if len(a) >= 4 and (a in needle or needle in a):
            return canonical
    return ai_payee.strip()


def _detect_nature_from_text(text: str) -> str | None:
    if not text:
        return None
    text_lower = text.lower()
    scores: dict = {}
    for nature, keywords in settings.NATURE_KEYWORDS.items():
        hits = sum(1 for kw in keywords if kw.lower() in text_lower)
        if hits > 0:
            scores[nature] = hits
    if not scores:
        return None
    return max(scores.items(), key=lambda kv: (kv[1], len(kv[0])))[0]


def _nature_for_payee(canonical_payee: str) -> str | None:
    preferred_order = settings.NATURE_PREFERRED_ORDER
    for nature in preferred_order:
        for comp_short, types in settings.COMPANY_TREE.items():
            for ptype, natures in types.items():
                if nature in natures and canonical_payee in natures[nature]:
                    return nature
    return None


def _find_nature_in_tree(nature_name: str) -> list:
    results = []
    for comp_short, types in settings.COMPANY_TREE.items():
        for ptype, natures in types.items():
            if nature_name in natures:
                results.append({
                    "company_short": comp_short,
                    "payment_type": ptype,
                    "payees": list(natures[nature_name]),
                })
    return results


def _regenerate_word(db: Session, req: Requisition) -> str:
    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req.id)
        .order_by(LineItem.item_number)
        .all()
    )

    company = None
    if req.company_id:
        company = db.query(Company).filter(Company.id == req.company_id).first()

    doc_data = {
        "requisition_id": str(req.id),
        "pr_number": req.pr_number,
        "applicant": req.applicant,
        "payee_name": req.payee_name,
        "nature": req.nature,
        "period": req.period,
        "category_name": req.category_name,
        "company_name": company.name if company else req.company_name,
        "company_address": company.address if company else None,
        "company_contact": None,
        "line_items": [
            {
                "date": li.date.strftime("%Y-%m-%d") if li.date else "",
                "description": li.description,
                "amount_hkd": li.amount_hkd,
                "category": li.category,
            }
            for li in line_items
        ],
    }
    return WordGenerator().generate_requisition_docx(doc_data)


def _generate_pr_number(db: Session, company_short: str | None = None) -> str:
    short = (company_short or "LSI").strip().upper()
    short = "".join(c for c in short if c.isalnum())[:6] or "LSI"

    year = datetime.utcnow().year
    prefix = f"{short}-PR-{year}-"

    rows = (
        db.query(Requisition)
        .filter(Requisition.pr_number.like(f"{prefix}%"))
        .all()
    )

    max_num = 0
    for r in rows:
        parts = (r.pr_number or "").split("-")
        if len(parts) >= 4:
            try:
                n = int(parts[-1])
                if n > max_num:
                    max_num = n
            except ValueError:
                continue

    return f"{prefix}{max_num + 1:03d}"


def _to_uuid(value):
    if value is None or value == "":
        return None
    if isinstance(value, UUID):
        return value
    try:
        return UUID(str(value))
    except (ValueError, AttributeError, TypeError):
        return None


def _infer_period(items) -> str | None:
    valid = []
    current_year = datetime.utcnow().year
    for item in items:
        d = _parse_date(item.get("date"))
        if not d:
            continue
        if d.year < current_year - 2 or d.year > current_year + 1:
            continue
        valid.append(d)
    if valid:
        return min(valid).strftime("%Y-%m")
    return None


def _parse_date(date_str):
    if not date_str:
        return None
    if isinstance(date_str, datetime):
        return date_str

    for fmt in (
        "%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y",
        "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
    ):
        try:
            return datetime.strptime(str(date_str).strip()[:19], fmt)
        except (ValueError, TypeError):
            continue
    return None


def _split_period(period: str | None) -> tuple:
    if not period or not isinstance(period, str):
        return ("—", "—")

    s = period.strip()

    if len(s) >= 7 and s[4] == "-":
        year = s[:4]
        month = s[5:7]
        if year.isdigit() and month.isdigit():
            return (year, month)

    if len(s) == 4 and s.isdigit():
        return (s, "—")

    return ("—", "—")


def _normalize_period(value):
    if not value:
        return None

    s = str(value).strip().replace("/", "-").replace(".", "-")

    import re as _re

    m = _re.match(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?", s)
    if m:
        y, mo = m.group(1), int(m.group(2))
        return f"{y}-{mo:02d}"

    if _re.match(r"^\d{4}$", s):
        return s

    return s or None


def _default_total_description(nature_name: str, doc_type: str | None = None) -> str:
    nature_lower = (nature_name or "").lower()
    doc_lower = (doc_type or "").lower()

    if "reimbursement" in nature_lower:
        return "Reimbursement total"
    if "mpf" in nature_lower:
        return "MPF contribution (total)"
    if "salary" in nature_lower:
        return "Salary payment"
    if "group medical" in nature_lower:
        return "Group Medical insurance premium"
    if "employment compensation" in nature_lower:
        return "Employment Compensation insurance"
    if "br renewal" in nature_lower:
        return "Business Registration renewal fee"
    if "auditor" in nature_lower:
        return "Auditor services"
    if "legal" in nature_lower:
        return "Legal consultation fee"
    if "company secretary" in nature_lower:
        return "Company Secretary fee"
    if "electricity" in nature_lower:
        return "Electricity bill"
    if "management fee" in nature_lower:
        return "Management fee"
    if "cleaning" in nature_lower:
        return "Office cleaning service"
    if "credit card" in nature_lower:
        return "Credit card payment"
    if "microsoft" in nature_lower:
        return "Microsoft 365 subscription"
    if "software" in nature_lower:
        return "Software annual fee"
    if doc_lower == "cheque":
        return "Cheque payment"
    return "Payment"